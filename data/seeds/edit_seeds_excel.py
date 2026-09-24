# -*- coding: utf-8 -*-
"""
把 data/seeds 下三份竞对清单 Excel 的公司名称规范化为唯一 canonical 名，
并把「官网链接」列的占位文本（如 "Abcam官网"）替换为真实 URL。
gen-site-template.xlsx 不参与。
"""
import openpyxl, glob, os

BASE = os.path.dirname(os.path.abspath(__file__))

# raw(单元格原值) -> (canonical_name, canonical_website)
MAP = {
    # ---- Elabscience 竞对清单 (Sheet2) ----
    "ABclonal": ("ABclonal", "https://www.abclonal.com/"),
    "ACROBiosystems": ("ACROBiosystems", "https://www.acrobiosystems.com/"),
    "Abcam": ("Abcam", "https://www.abcam.com/"),
    "BD Biosciences": ("BD Biosciences", "https://www.bdbiosciences.com/"),
    "Beckman Coulter": ("Beckman Coulter", "https://www.beckman.com/"),
    "Beyotime / 碧云天": ("碧云天（Beyotime）", "https://www.beyotime.com/"),
    "Bio X Cell": ("Bio X Cell", "https://www.bioxcell.com/"),
    "BioAssay Systems": ("BioAssay Systems", "https://www.bioassaysystems.com/"),
    "BioLegend": ("BioLegend", "https://www.biolegend.com/"),
    "CST": ("CST（Cell Signaling Technology）", "https://www.cellsignal.com/"),
    "CST / Cell Signaling Technology": ("CST（Cell Signaling Technology）", "https://www.cellsignal.com/"),
    "Cayman Chemical": ("Cayman Chemical", "https://www.caymanchem.com/"),
    "Leinco": ("Leinco", "https://www.leinco.com/"),
    "MedChemExpress / MCE": ("MedChemExpress（MCE）", "https://www.medchemexpress.com/"),
    "Miltenyi Biotec": ("Miltenyi Biotec", "https://www.miltenyibiotec.com/"),
    "Promega": ("Promega", "https://www.promega.com/"),
    "Proteintech": ("Proteintech", "https://www.ptglab.com/"),
    "R&D Systems": ("R&D Systems", "https://www.rndsystems.com/"),
    "RayBiotech": ("RayBiotech", "https://www.raybiotech.com/"),
    "STEMCELL Technologies": ("STEMCELL Technologies", "https://www.stemcell.com/"),
    "Selleck Chemicals": ("Selleck", "https://www.selleckchem.com/"),
    "Sino Biological / 义翘神州": ("义翘神州（Sino Biological）", "https://www.sinobiological.com/"),
    "Solarbio / 索莱宝": ("索莱宝（Solarbio）", "https://www.solarbio.com/"),
    "Thermo Fisher / Invitrogen": ("赛默飞（Thermo Fisher Scientific）", "https://www.thermofisher.com/"),
    "eBioscience / Thermo Fisher": ("赛默飞（Thermo Fisher Scientific）", "https://www.thermofisher.com/"),
    "华安生物 / HUABIO": ("华安生物（HUABIO）", "https://www.huabio.com/"),
    "华美生物 / CUSABIO": ("华美生物（CUSABIO）", "https://www.cusabio.com/"),
    "同仁化学 / Dojindo": ("同仁化学（Dojindo）", "https://www.dojindo.com/"),
    "同立海源": ("同立海源", "https://www.seafrom.com/"),
    "四正柏 / 4A Biotech": ("四正柏（4A Biotech）", "https://www.4abiotech.com/"),
    "斯达特 / Starter": ("斯达特（Starter）", "https://www.starter-bio.com/"),
    "格锐思 / Geruisi": ("格锐思（Geruisi）", "https://www.geruisi-bio.com/"),
    "江莱生物 / Jonln": ("江莱生物（Jonln）", "http://www.jonln.com/"),
    "海狸 / Beaver": ("海狸（Beaver）", "https://www.beaverbio.com/"),
    "盒子生工 / Boxbio": ("盒子生工（Boxbio）", "https://www.boxbio.cn/"),
    "联科生物 / Liankebio": ("联科生物（Liankebio）", "https://www.liankebio.com/"),
    "菲恩生物 / FineTest": ("菲恩生物（FineTest）", "https://www.fn-test.com/"),
    "近岸蛋白 / Novoprotein": ("近岸蛋白（Novoprotein）", "https://www.novoprotein.com.cn/"),
    "金斯瑞 / GenScript": ("金斯瑞（GenScript）", "https://www.genscript.com/"),
    # ---- Procell 竞对清单 ----
    "美国ATCC": ("ATCC", "https://www.atcc.org/"),
    "德国promocell": ("PromoCell", "https://www.promocell.com/"),
    "德国cytion": ("Cytion", "https://www.cytion.com/"),
    "上海中乔新舟": ("上海中乔新舟生物科技有限公司", "https://www.zqxzbio.com/"),
    "上海赛百慷": ("赛百慷（上海）生物技术股份有限公司", "https://www.icellbioscience.com/"),
    "上海富衡": ("上海富衡生物科技有限公司", "https://www.fudancell.com/"),
    "苏州海星": ("苏州海星生物科技有限公司", "https://www.cas9x.com/"),
    "浙江美森": ("浙江美森细胞科技有限公司", "https://ctcc.online/"),
    "广州赛业": ("赛业（广州）生物科技有限公司", "https://www.oricellbio.cn/"),
    "厦门逸漠": ("厦门逸漠生物科技有限公司", "http://immocell.com/"),
    "赛默飞": ("赛默飞（Thermo Fisher Scientific）", "https://www.thermofisher.cn/"),
    "stemcell": ("STEMCELL Technologies", "https://www.stemcell.com/"),
    "sciencell": ("ScienCell Research Laboratories", "https://www.sciencellonline.com/"),
    "德国Capricorn": ("Capricorn Scientific", "https://www.capricorn-scientific.com/"),
    "默克": ("默克（Merck）", "https://www.sigmaaldrich.com/"),
    "康宁": ("康宁（Corning）", "https://www.corning.com/"),
    "上海逍鹏": ("上海逍鹏生物科技有限公司", "https://www.xpbiomed.com/"),
    "上海源培": ("上海源培生物科技股份有限公司", "https://www.basalmedia.com/"),
    "依科赛": ("苏州依科赛生物科技股份有限公司", "https://www.excellbio.com/"),
    "南京森贝伽": ("南京森贝伽生物科技有限公司", "http://www.senbeijia.com/"),
    "苏州双洳": ("苏州双洳生物科技有限公司", "https://www.lonsera.cn/"),
    "诺唯赞": ("南京诺唯赞生物科技股份有限公司", "https://www.vazyme.com/"),
    "cytiva(HyClone)": ("Cytiva", "https://www.cytivalifesciences.com/"),
    # ---- 补充-竞品清单 ----
    "逍鹏生物": ("上海逍鹏生物科技有限公司", "https://www.xpbiomed.com/"),
}

TARGETS = ["Elabscience竞对清单.xlsx", "Procell 竞对清单-2026.9.xlsx", "补充-竞品清单.xlsx"]

def colidx(header, *keys):
    for i, h in enumerate(header):
        if h and any(k in str(h) for k in keys):
            return i
    return None

total = 0
for fn in TARGETS:
    path = os.path.join(BASE, fn)
    if not os.path.exists(path):
        print("SKIP (missing):", fn); continue
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        header = list(rows[0])
        ci = colidx(header, "公司", "竞对")
        wi = colidx(header, "官网", "域名", "网址")
        if ci is None:
            continue
        changed = 0
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=ci + 1)
            raw = str(cell.value).strip() if cell.value else ""
            if raw in MAP:
                canon, web = MAP[raw]
                cell.value = canon
                if wi is not None:
                    wc = ws.cell(row=r, column=wi + 1)
                    wc.value = web
                changed += 1
        total += changed
        print(f"  {fn} / {ws.title}: 规范化 {changed} 行")
    wb.save(path)
    print("  saved:", fn)

print("DONE. 共规范化", total, "行")
